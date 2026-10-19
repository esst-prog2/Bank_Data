"""Unit tests for reconcile.py's own parsing logic, using hand-built, frozen inputs -
no cached download, no network, nothing that can skip depending on what's been
fetched. This is the gap homework 6 review 8 found: every test in test_reconcile.py
exercises reconcile.py through real cached files under data/raw/ (an integration
test - checks the pieces connect), and 10 of them `pytest.skip()` when that cache is
missing. Nothing pinned down reconcile.py's own processing logic in isolation (a unit
test - checks the logic itself), so a fresh clone with no cached data runs almost none
of the suite and still goes green.

This file also carries a permanent regression guard for a real bug (see
AGENTS.md/PLANNING_LOG.md, 2026-10-12): openpyxl's `read_only=True` `iter_rows()`
trims a row's tuple to its last non-empty cell, so a row whose trailing columns are
genuinely blank is SHORTER than the header - the original code's plain `row[i]`
crashed with IndexError on exactly this shape (first found on BNP Paribas's real LI1
export). `reconcile._read_p3dh_fact`'s bounds-checking logic was extracted to a
module-level `_row_cell()` (previously a nested closure) specifically so this could
be tested with a hand-built short tuple directly, rather than only relying on
whichever bank's cached export happens to still have a blank trailing cell today -
or on openpyxl's own writer happening to reproduce the same trimming (it doesn't;
confirmed while writing this test - a file built via `ws.append()` pads short rows
with `None` instead of truly trimming them).
"""
from __future__ import annotations

from datetime import datetime

import openpyxl
import pandas as pd
import pytest

from data_acquisition import reconcile


def _write_p3dh_like_export(path, header, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([datetime(2025, 12, 31)])  # row 1: ReferenceDate, like a real P3DH export
    ws.append(header)  # row 2: the real header
    for row in rows:
        ws.append(row)  # a short `row` leaves trailing cells genuinely untouched
    wb.save(path)


_HEADER = ["Row", "Open Key", "Column", "Column Name", "Row Name", "FactValue"]


def test_read_p3dh_fact_returns_value_for_a_fully_populated_row(tmp_path):
    path = tmp_path / "fake_export.xlsx"
    _write_p3dh_like_export(path, _HEADER, [
        ["0010", None, "0010", "a. T", "CET1 capital", 1234.5],
    ])

    reference_date, description, column_desc, value = reconcile._read_p3dh_fact(
        path, "Row", "0010", "0010"
    )

    assert reference_date == datetime(2025, 12, 31)
    assert description == "CET1 capital"
    assert column_desc == "a. T"
    assert value == 1234.5


def test_row_cell_reads_a_row_trimmed_shorter_than_the_header_as_blank():
    """The exact BNP Paribas LI1 shape: openpyxl's read_only iter_rows() hands back a
    row tuple shorter than the header whenever that row's trailing cells are
    genuinely blank - a plain row[i] raises IndexError on those columns instead of
    reading them as blank (the real bug, fixed 2026-10-12). Hand-built tuple, not a
    file - this is the bounds-checking logic itself, isolated from whether openpyxl's
    own writer happens to reproduce the same trimming today."""
    idx = {name: i for i, name in enumerate(_HEADER)}
    short_row = ("0020", None, "0010", "a. T", "Some other row")  # 5 cells, not 6

    assert reconcile._row_cell(short_row, idx, "Row Name") == "Some other row"
    assert reconcile._row_cell(short_row, idx, "FactValue") is None  # would IndexError without the fix


def test_row_cell_reads_a_full_length_row_normally():
    idx = {name: i for i, name in enumerate(_HEADER)}
    full_row = ("0010", None, "0010", "a. T", "CET1 capital", 1234.5)

    assert reconcile._row_cell(full_row, idx, "FactValue") == 1234.5


def test_read_p3dh_fact_raises_on_malformed_export_missing_reference_date(tmp_path):
    path = tmp_path / "fake_export.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["not a date"])  # row 1 has no datetime anywhere
    ws.append(_HEADER)
    wb.save(path)

    with pytest.raises(reconcile.DataSourceError):
        reconcile._read_p3dh_fact(path, "Row", "0010", "0010")


@pytest.mark.parametrize("period, expected", [("202412", "2024-12-31"), ("202112", "2021-12-31")])
def test_period_to_esef_period_end_round_trips(period, expected):
    period_end = reconcile._period_to_esef_period_end(period)
    assert period_end == expected
    assert reconcile._esef_period_end_to_period(period_end) == period


def test_period_to_esef_period_end_rejects_non_december_year_end():
    with pytest.raises(ValueError):
        reconcile._period_to_esef_period_end("202409")


# Homework 6 review feedback (2026-10-19): test_reconcile.py's own compute_npl_ratio/
# compute_dupont tests either check the function against a value computed by RUNNING
# the function itself (no independent source exists, so a wrong CQ3 mapping would
# simply agree with itself), or check an algebraic identity that holds for ANY inputs
# (ROE = Profit Margin x Asset Utilization x Equity Multiplier reduces to Net Income /
# Total Equity by construction). Neither can catch a wrong formula or a wrong mapping
# against a value an outsider could check. These monkeypatch get_financial_data with
# small, hand-built, independently-known inputs so the expected outputs are computed
# on paper, not by running the code under test.

def test_compute_npl_ratio_hand_calculation(monkeypatch):
    def fake_get_financial_data(bank_lei, concepts, period=None, csv_path=None):
        return pd.DataFrame([
            {"standardized_concept": "performing_exposures", "status": "ok", "value": 980.0},
            {"standardized_concept": "non_performing_exposures", "status": "ok", "value": 20.0},
        ])

    monkeypatch.setattr(reconcile, "get_financial_data", fake_get_financial_data)

    ratio = reconcile.compute_npl_ratio("ANY_LEI", "202512")
    assert ratio == pytest.approx(0.02)  # 20 / (980 + 20), by hand


def test_compute_dupont_hand_calculation(monkeypatch):
    # Hand-picked so every DuPont factor is a distinct, checkable-on-paper number:
    # total_revenue = 200 (NII) + 50 (net fee) + 0 (trading) = 250.
    values = {
        "net_interest_income": 200.0,
        "fee_and_commission_income_expense": 50.0,
        "trading_income_expense": 0.0,
        "profit_or_loss_for_the_year": 100.0,
        "total_assets": 2000.0,
        "total_equity": 200.0,
    }

    def fake_get_financial_data(bank_lei, concepts, period=None, csv_path=None):
        return pd.DataFrame([
            {"standardized_concept": c, "source": "esef", "status": "ok", "value": v}
            for c, v in values.items()
        ])

    monkeypatch.setattr(reconcile, "get_financial_data", fake_get_financial_data)

    result = reconcile.compute_dupont("ANY_LEI", "202412")

    assert result.total_revenue == pytest.approx(250.0)
    assert result.profit_margin == pytest.approx(100.0 / 250.0)  # 0.4
    assert result.asset_utilization == pytest.approx(250.0 / 2000.0)  # 0.125
    assert result.equity_multiplier == pytest.approx(2000.0 / 200.0)  # 10.0
    assert result.roa == pytest.approx(0.05)  # 100 / 2000
    assert result.roe == pytest.approx(0.5)  # 100 / 200, and 0.4 * 0.125 * 10.0
